"""Real auth/chat/stream regression. Sends stop at confirmation; no confirmations are submitted."""
from __future__ import annotations
import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import sys
import requests
from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from app.db import open_pool, close_pool, fetch_one
from app.services.context_service import get_thread_state

BASE = os.getenv('VERIFY_API_BASE_URL', 'http://127.0.0.1:8002').rstrip('/')
ACCOUNTS = {'admin':('admin_demo','Admin123456'), 'finance':('finance_demo','Finance123456'),
            'operations':('operations_demo','Operations123456'), 'customer_service':('employee_demo','Employee123456')}

def login(role):
    response = requests.post(BASE+'/auth/login', data=dict(zip(('username','password'),ACCOUNTS[role])), timeout=30)
    response.raise_for_status()
    return response.json()['access_token']

def call(token, endpoint, message, thread_id=None):
    response = requests.post(BASE+endpoint, headers={'Authorization':'Bearer '+token},
        json={'message': message, **({'thread_id':thread_id} if thread_id else {})}, timeout=240, stream=endpoint.endswith('stream'))
    if response.status_code != 200:
        raise AssertionError((endpoint, message, response.status_code, response.text[:500]))
    if not endpoint.endswith('stream'):
        return response.json()
    event = ''
    for line in response.iter_lines(decode_unicode=True):
        if line.startswith('event:'):
            event=line[6:].strip()
        if line.startswith('data:'):
            value=json.loads(line[5:].strip())
            if event=='error':
                raise AssertionError(value)
            if event=='done':
                return value
    raise AssertionError('SSE missing done')

def compact(result):
    approval = result.get('approval_result') or {}
    automation = result.get('automation') or {}
    return {'intent':result.get('intent'),'answer':result.get('answer','')[:450], 'thread_id':result.get('thread_id'),
            'attachment_count':len(result.get('attachments') or []),'citation_count':len(result.get('citations') or automation.get('citations') or []),
            'status':approval.get('status') or automation.get('status'), 'run_id':automation.get('run_id')}

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--output',type=Path,default=ROOT/'data/chat-intent-eval/api-round1.json')
    args=parser.parse_args()
    tokens={role:login(role) for role in ACCOUNTS}
    cases=[('leave','admin','请假找谁请','policy'), ('process','finance','请假流程是什么','policy'),
           ('guide','finance','工资表应该发给谁？','policy'),('hello','operations','你好','chitchat'),
           ('thanks','customer_service','谢谢','chitchat'),('unsupported','admin','帮我提交请假申请','ask_clarification'),
           ('ambiguous','finance','导出excel表','ask_clarification'),('negated','admin','不要发送邮件','ask_clarification'),
           ('denied','finance','帮我生成listing文案并保存草稿','permission_denied'),
           ('salary','finance','导出2026年07月所有员工工资表Excel','finance_salary_export'),
           ('salary-polite','finance','能把2026年07月工资单导出来吗','finance_salary_export'),
           ('salary-capability','finance','系统能导出工资单吗','policy'),
           ('compound-missing','finance','帮我生成这个月的财务报表和工资表','ask_clarification'),
           ('compound-files','finance','生成2026年07月财务报表和工资表','finance_compound_report_generation'),
           ('refund-rule','finance','公司退款规定是什么','policy')]
    def run_case(case):
        name,role,message,expected=case
        pair=[]
        for endpoint in ['/chat','/chat/stream']:
            result=call(tokens[role],endpoint,message)
            assert result.get('intent')==expected, (name,endpoint,compact(result))
            if expected=='policy':
                assert not result.get('automation') and not result.get('approval_result') and not result.get('attachments'),compact(result)
                slots=get_thread_state(result['thread_id'])['slots']
                assert not (slots.get('pending_external_action') or {}).get('active'),slots
                assert not (slots.get('execution_context') or {}).get('active'),slots
            if expected in {'finance_salary_export','finance_compound_report_generation'}:
                assert result.get('attachments'),compact(result)
                if expected=='finance_compound_report_generation':
                    assert len(result['attachments'])>=2,compact(result)
                item=result['attachments'][0]
                path=item.get('download_path') or (item.get('metadata') or {}).get('download_path')
                assert path, {k:v for k,v in item.items() if k!='content_base64'}
                file=requests.get(BASE+path,headers={'Authorization':'Bearer '+tokens[role]},timeout=30)
                file.raise_for_status()
                workbook=load_workbook(BytesIO(file.content),read_only=True)
                assert workbook.active.max_row>=2
            pair.append(compact(result))
        print(json.dumps({'case':name,'intents':[r['intent'] for r in pair]},ensure_ascii=False),flush=True)
        return {'id':name,'role':role,'expected':expected,'pair':pair,'consistent':pair[0]['intent']==pair[1]['intent']}
    open_pool()
    try:
        with ThreadPoolExecutor(max_workers=3) as pool:
            results=list(pool.map(run_case,cases))
        for endpoint in ['/chat','/chat/stream']:
            # Create a real pending generic file send. New question must close it.
            initial=call(tokens['admin'],endpoint,'把刚生成的文件发到邮箱')
            assert initial.get('intent')=='chat_plan_execute',compact(initial)
            state=get_thread_state(initial['thread_id'])
            assert state['slots']['pending_external_action']['active'],compact(initial)
            question=call(tokens['admin'],endpoint,'请假找谁请',initial['thread_id'])
            assert question.get('intent')=='policy' and not question.get('approval_result'),compact(question)
            state=get_thread_state(question['thread_id'])
            assert not state['slots']['pending_external_action']['active'],state
            results.append({'id':'pending-switch','endpoint':endpoint,'pair':[compact(initial),compact(question)]})
            print(json.dumps({'case':'pending-switch','endpoint':endpoint,'intent':question.get('intent')},ensure_ascii=False),flush=True)
            # Valid email remains a continuation and reaches attachment lookup.
            initial=call(tokens['admin'],endpoint,'把刚生成的文件发到邮箱')
            followup=call(tokens['admin'],endpoint,'intent-routing-test@example.com',initial['thread_id'])
            assert followup.get('intent')=='chat_plan_execute',compact(followup)
            assert not (followup.get('approval_result') or {}).get('sent'),compact(followup)
            results.append({'id':'valid-continuation','endpoint':endpoint,'pair':[compact(initial),compact(followup)]})
            print(json.dumps({'case':'valid-continuation','endpoint':endpoint,'intent':followup.get('intent')},ensure_ascii=False),flush=True)
        # Knowledge-first compound request must never send an older attachment.
        for endpoint in ['/chat','/chat/stream']:
            compound=call(tokens['admin'],endpoint,'先查公司退款规定，再把结果发送到 intent-routing-test@example.com')
            assert compound.get('intent')=='chat_plan_execute',compact(compound)
            assert compound.get('attachments'),compact(compound)
            confirmation=compound.get('approval_result') or {}
            assert confirmation.get('status')=='waiting_email_confirmation',compact(compound)
            assert confirmation.get('requires_confirmation') is True,confirmation
            assert confirmation['artifact']['filename']=='知识查询结果.txt',confirmation
            assert 'intent-routing-test@example.com'==confirmation.get('recipient'),confirmation
            results.append({'id':'knowledge-then-send','endpoint':endpoint,'result':compact(compound)})
            print(json.dumps({'case':'knowledge-then-send','endpoint':endpoint,'status':confirmation.get('status')},ensure_ascii=False),flush=True)
        for endpoint in ['/chat','/chat/stream']:
            for message,expected in [('你好','chitchat'),('谢谢','chitchat')]:
                initial=call(tokens['admin'],endpoint,'把刚生成的文件发到邮箱')
                answer=call(tokens['admin'],endpoint,message,initial['thread_id'])
                assert answer.get('intent')==expected and not answer.get('approval_result'),compact(answer)
                state=get_thread_state(answer['thread_id'])
                assert not state['slots']['pending_external_action']['active'],state
                results.append({'id':'pending-'+message,'endpoint':endpoint,'result':compact(answer)})
                print(json.dumps({'case':'pending-'+message,'endpoint':endpoint,'intent':answer.get('intent')},ensure_ascii=False),flush=True)
        report={'timestamp':datetime.now(timezone.utc).isoformat(),'base':BASE,'mode':'real_api_real_auth_real_model_db',
                'result_count':len(results),'paired_case_count':len(cases),'consistent_pairs':sum(r.get('consistent',False) for r in results),'results':results,
                'note':'No external send confirmation was submitted. Test threads/artifacts remain for trace inspection.'}
        args.output.parent.mkdir(parents=True,exist_ok=True)
        args.output.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
        print(json.dumps({k:v for k,v in report.items() if k!='results'},ensure_ascii=False,indent=2))
    finally:
        close_pool()

if __name__=='__main__':
    main()
