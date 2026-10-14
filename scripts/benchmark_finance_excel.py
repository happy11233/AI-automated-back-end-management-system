from __future__ import annotations

import argparse
from datetime import datetime, timezone
from io import BytesIO
import json
import os
from pathlib import Path
import statistics
import time
from typing import Any

import pandas as pd
import requests
from openpyxl import load_workbook


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    if not ordered:
        return 0.0
    index = min(len(ordered) - 1, max(0, int((len(ordered) - 1) * fraction)))
    return ordered[index]


def main() -> None:
    args = parse_args()
    base_url = args.api_base_url.rstrip("/")
    token = login(base_url, args.username, args.password)
    headers = {"Authorization": f"Bearer {token}"}
    content = workbook_bytes(args.order_prefix)
    samples: list[dict[str, Any]] = []

    for index in range(1, args.count + 1):
        started = time.perf_counter()
        response = requests.post(
            f"{base_url}/automation/finance/excel-transform",
            headers=headers,
            files={"file": ("finance_latency_benchmark.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"instruction": args.instruction, "erp_resources": "[]"},
            timeout=args.timeout,
        )
        duration_ms = round((time.perf_counter() - started) * 1000)
        response.raise_for_status()
        if response.content[:2] != b"PK":
            raise RuntimeError(f"sample {index} did not return an xlsx file")
        workbook = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        samples.append({
            "index": index,
            "duration_ms": duration_ms,
            "response_bytes": len(response.content),
            "sheet_count": len(workbook.sheetnames),
            "has_summary_sheet": "处理摘要" in workbook.sheetnames,
        })
        print(json.dumps({"sample": index, "duration_ms": duration_ms, "response_bytes": len(response.content)}, ensure_ascii=False), flush=True)

    durations = [sample["duration_ms"] / 1000 for sample in samples]
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "run_label": args.run_label,
        "scenario": "finance_excel_pure_formatting_fixed_input",
        "configuration": {
            "api_base_url": base_url,
            "instruction": args.instruction,
            "count": args.count,
            "input_bytes": len(content),
        },
        "summary": {
            "samples": len(durations),
            "successes": sum(sample["has_summary_sheet"] for sample in samples),
            "avg_seconds": round(statistics.mean(durations), 3) if durations else 0,
            "p50_seconds": round(percentile(durations, 0.50), 3),
            "p95_seconds": round(percentile(durations, 0.95), 3),
            "p99_seconds": round(percentile(durations, 0.99), 3),
            "max_seconds": round(max(durations), 3) if durations else 0,
        },
        "samples_detail": samples,
    }
    output = Path(args.output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"ok": True, "output": str(output), "summary": report["summary"]}, ensure_ascii=False))


def workbook_bytes(order_prefix: str) -> bytes:
    output = BytesIO()
    frame = pd.DataFrame([
        {"posting_date": "2026-09-01", "gross_pay": 10000, "employee_name": "Benchmark One", "currency": "USD"},
        {"posting_date": "2026-09-02", "gross_pay": 12000, "employee_name": "Benchmark Two", "currency": "USD"},
    ])
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False, sheet_name="Monthly Report")
    return output.getvalue()


def login(base_url: str, username: str, password: str) -> str:
    response = requests.post(f"{base_url}/auth/login", data={"username": username, "password": password}, timeout=30)
    response.raise_for_status()
    return response.json()["access_token"]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run a fixed-input real finance Excel formatting benchmark")
    parser.add_argument("--count", type=int, default=30)
    parser.add_argument("--run-label", default="baseline-optimized")
    parser.add_argument("--api-base-url", default=os.getenv("VERIFY_API_BASE_URL", "http://127.0.0.1:8001"))
    parser.add_argument("--username", default="finance_demo")
    parser.add_argument("--password", default="Finance123456")
    parser.add_argument("--order-prefix", default="FINANCE-BENCHMARK")
    parser.add_argument("--instruction", default="将财务报表里面的内容给美化一下，每个字段都是英文单词不好读")
    parser.add_argument("--timeout", type=int, default=120)
    parser.add_argument("--output", default="data/latency-benchmarks/finance-excel-baseline-optimized.json")
    return parser.parse_args()


if __name__ == "__main__":
    main()
