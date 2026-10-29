"""Verify financial quality rules with real XLSX input/output and no test doubles.

All amounts are isolated verification data, not actual company transactions.
This script verifies deterministic file processing, not API/platform integration.
"""
from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path
import sys
import unittest

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.services.finance_reconciliation_service import (  # noqa: E402
    FinanceReconciliationInputFile,
    reconcile_finance_workbooks,
)


def input_file(filename, headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return FinanceReconciliationInputFile(filename, output.getvalue())


def fixtures(*, currency="USD", sales=100, rate=7, purchase=True, logistics=True):
    files = [input_file("settlement.xlsx",
                        ["订单号", "SKU", "数量", "币种", "销售额", "退款", "平台手续费"],
                        [["VERIFY-FIN-001", "VERIFY-SKU-1", 2, currency, sales, 0, 10]])]
    if purchase:
        files.append(input_file("purchase.xlsx", ["SKU", "币种", "采购成本"],
                                [["VERIFY-SKU-1", "CNY", 100]]))
    if logistics:
        files.append(input_file("logistics.xlsx", ["订单号", "币种", "物流费"],
                                [["VERIFY-FIN-001", "CNY", 30]]))
    if rate is not None:
        files.append(input_file("exchange_rate.xlsx", ["币种", "汇率"], [["USD", rate]]))
    return files


def inspect_result(files):
    result = reconcile_finance_workbooks(files=files)
    workbook = load_workbook(BytesIO(result.content), data_only=True)
    values = list(workbook["订单利润表"].values)
    lines = [dict(zip(values[0], row)) for row in values[1:]]
    summary = dict(list(workbook["对账摘要"].values)[1:])
    json.dumps(result.metadata, allow_nan=False)
    return result, lines, summary


class FinanceQualityVerification(unittest.TestCase):
    def test_complete_inputs_have_traceable_contribution_profit(self):
        result, lines, summary = inspect_result(fixtures())
        self.assertEqual(lines[0]["销售额"], 700)
        self.assertEqual(lines[0]["采购成本"], 200)
        self.assertEqual(lines[0]["利润"], 400)
        self.assertAlmostEqual(lines[0]["利润率"], 400 / 700)
        self.assertEqual(result.metadata["total_profit"], 400)
        self.assertEqual(result.metadata["review_status"], "pending_review")
        self.assertIn("银行到账", summary["报告口径"])
        self.assertIn("settlement.xlsx", lines[0]["来源"])

    def test_missing_settlement_exchange_rate_requires_input(self):
        with self.assertRaisesRegex(ValueError, "缺少 USD.*汇率"):
            inspect_result(fixtures(rate=None))

    def test_missing_exchange_rate_in_every_cost_source_requires_input(self):
        for filename, headers, row in [
            ("purchase_eur.xlsx", ["SKU", "币种", "采购成本"], ["VERIFY-SKU-1", "EUR", 10]),
            ("logistics_eur.xlsx", ["订单号", "币种", "物流费"], ["VERIFY-FIN-001", "EUR", 10]),
            ("ads_eur.xlsx", ["SKU", "币种", "广告费"], ["VERIFY-SKU-1", "EUR", 10]),
        ]:
            with self.subTest(source=filename), self.assertRaisesRegex(ValueError, "缺少 EUR.*汇率"):
                inspect_result(fixtures() + [input_file(filename, headers, [row])])

    def test_base_currency_needs_no_exchange_file(self):
        result, lines, _ = inspect_result(fixtures(currency="CNY", rate=None, sales=500))
        self.assertEqual(lines[0]["利润"], 260)
        self.assertEqual(result.metadata["total_profit"], 260)

    def test_blank_currency_in_mixed_currency_file_requires_input(self):
        extra = input_file("mixed_settlement.xlsx", ["订单号", "SKU", "币种", "销售额"],
                           [["VERIFY-FIN-002", "VERIFY-SKU-2", "CNY", 10],
                            ["VERIFY-FIN-003", "VERIFY-SKU-3", "", 20]])
        with self.assertRaisesRegex(ValueError, "币种列存在空值"):
            inspect_result(fixtures() + [extra])

    def test_actual_one_to_one_foreign_rate_is_allowed(self):
        result, lines, _ = inspect_result(fixtures(rate=1))
        self.assertEqual(lines[0]["销售额"], 100)
        self.assertEqual(result.metadata["total_profit"], -140)

    def test_invalid_exchange_rates_require_correction(self):
        for rate in (0, -2, "unknown", "", "Infinity"):
            with self.subTest(rate=rate), self.assertRaisesRegex(ValueError, "无效币种或汇率|缺少 USD.*汇率"):
                inspect_result(fixtures(rate=rate))

    def test_conflicting_exchange_rates_are_not_silently_overwritten(self):
        for currency, rate in (("USD", 8), ("CNY", 7)):
            with self.subTest(currency=currency), self.assertRaisesRegex(ValueError, "冲突汇率"):
                inspect_result(fixtures() + [input_file("extra_rate.xlsx", ["币种", "汇率"], [[currency, rate]])])

    def test_identical_exchange_rates_can_be_repeated(self):
        result, _, _ = inspect_result(fixtures() + [input_file("extra_rate.xlsx", ["币种", "汇率"], [["USD", 7]])])
        self.assertEqual(result.metadata["total_profit"], 400)

    def test_missing_costs_do_not_produce_profit_or_margin(self):
        for field in ("purchase", "logistics"):
            with self.subTest(field=field):
                result, lines, summary = inspect_result(fixtures(**{field: False}))
                self.assertIn("待补", lines[0]["状态"])
                for key in ("总成本", "利润", "利润率"):
                    self.assertIsNone(lines[0][key])
                self.assertIsNone(result.metadata["total_profit"])
                self.assertIsNone(result.metadata["calculable_profit_subtotal"])
                self.assertEqual(result.metadata["incomplete_line_count"], 1)
                self.assertEqual(result.metadata["calculable_line_count"], 0)
                self.assertEqual(result.metadata["review_status"], "requires_input")
                self.assertIsNone(summary["利润合计"])
                self.assertEqual(summary["复核状态"], "待补成本资料")

    def test_missing_cost_and_zero_sales_do_not_produce_zero_margin(self):
        _, lines, _ = inspect_result(fixtures(sales=0, purchase=False))
        self.assertIsNone(lines[0]["利润率"])

    def test_partial_subtotal_never_appears_as_total_profit(self):
        extra = input_file("extra_settlement.xlsx", ["订单号", "SKU", "币种", "销售额"],
                           [["VERIFY-FIN-002", "VERIFY-SKU-2", "CNY", 1000]])
        result, lines, summary = inspect_result(fixtures() + [extra])
        self.assertEqual(result.metadata["calculable_profit_subtotal"], 400)
        self.assertEqual(result.metadata["calculable_line_count"], 1)
        self.assertEqual(result.metadata["incomplete_line_count"], 1)
        self.assertIsNone(result.metadata["total_profit"])
        self.assertIsNone(summary["利润合计"])
        self.assertEqual(summary["可计算行利润小计"], 400)
        self.assertIsNone(next(row for row in lines if row["订单号"] == "VERIFY-FIN-002")["利润"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
